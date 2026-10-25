# -*- coding: utf-8 -*-
"""
Мастер Веб-Апликација (Десктоп + Мобилен) за ФБ Контрола,
Длабински Фузи-Пресеци, Преглед на Линкови од Објави,
Управување со Тимови (Додавање/Бришење/Замена) и Телеграм Потсетници во Две Брзини.
"""
import streamlit as st
import streamlit.components.v1 as components
import pandas as pd
import numpy as np
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import os
import re
import io
import json
import socket
import urllib.parse
from datetime import datetime
from difflib import SequenceMatcher

st.set_page_config(
    page_title="ФБ Контрола - Мастер Систем",
    page_icon="📊",
    layout="wide",
    initial_sidebar_state="expanded"
)

# Стилизирање за Десктоп и Мобилен приказ
st.markdown("""
<style>
    .main-header { font-size: 26px; font-weight: 800; color: #1F497D; margin-bottom: 6px; }
    .sub-header { font-size: 19px; font-weight: 700; color: #274E13; margin-top: 12px; margin-bottom: 8px; }
    .link-card {
        background: linear-gradient(135deg, #f0f4ff 0%, #e4ecfa 100%);
        border: 2px solid #1F497D;
        border-radius: 10px;
        padding: 14px 18px;
        margin-bottom: 14px;
    }
    .link-badge {
        background-color: #1F497D;
        color: white;
        padding: 4px 10px;
        border-radius: 6px;
        font-size: 13px;
        font-weight: bold;
    }
    .warn-badge {
        background-color: #990000;
        color: white;
        padding: 4px 10px;
        border-radius: 6px;
        font-size: 13px;
        font-weight: bold;
    }
</style>
""", unsafe_allow_html=True)

# ==============================================================================
# 0. ЗАШТИТА СО ЛОЗИНКА (LOGIN SYSTEM)
# ==============================================================================
# Овде можеш да ги смениш корисничките имиња и лозинките
USER_CREDENTIALS = {
    "admin": "admin123",
    "toni": "skopje2026"
}

if "authenticated" not in st.session_state:
    st.session_state["authenticated"] = False

def logout():
    st.session_state["authenticated"] = False

if not st.session_state["authenticated"]:
    st.markdown('<div class="main-header">🔒 Безбедносна Најава</div>', unsafe_allow_html=True)
    st.info("Апликацијата е заштитена. Внесете корисничко име и лозинка за пристап.")
    
    with st.form("login_form"):
        username = st.text_input("Корисничко име:")
        password = st.text_input("Лозинка:", type="password")
        submit_btn = st.form_submit_button("Најави се")
        
        if submit_btn:
            if username in USER_CREDENTIALS and USER_CREDENTIALS[username] == password:
                st.session_state["authenticated"] = True
                st.success("Успешна најава! Се вчитува системот...")
                st.rerun()
            else:
                st.error("❌ Погрешно корисничко име или лозинка!")
    st.stop() # Го стопира вчитувањето на остатокот од апликацијата ако не е најавен

# Додавање копче за одјава во левото мени (се појавува само кога е најавен)
st.sidebar.button("🚪 Одјави се", on_click=logout)
st.sidebar.markdown("---")


# ==============================================================================
# 1. КОМПЛЕТЕН АЛИЈАС РЕЧНИК И НОРМАЛИЗАЦИЈА (Сите одлуки од разговорите)
# ==============================================================================
CYR_TO_LAT = {
    'а': 'a', 'б': 'b', 'в': 'v', 'г': 'g', 'д': 'd', 'ѓ': 'gj', 'е': 'e', 'ж': 'z',
    'з': 'z', 'ѕ': 'dz', 'и': 'i', 'ј': 'j', 'к': 'k', 'л': 'l', 'љ': 'lj', 'м': 'm',
    'н': 'n', 'њ': 'nj', 'о': 'o', 'п': 'p', 'р': 'r', 'с': 's', 'т': 't', 'ќ': 'k',
    'у': 'u', 'ф': 'f', 'х': 'h', 'ц': 'c', 'ч': 'c', 'џ': 'dz', 'ш': 's',
    'ć': 'c', 'č': 'c', 'š': 's', 'ž': 'z', 'đ': 'dj', 'ћ': 'c'
}

def normalize_name(s):
    if not isinstance(s, str):
        return ""
    s = s.lower()
    s = s.split(' (http')[0].split(' http')[0]
    s = s.replace('[', '').replace(']', '').replace(' - предложувач', '').replace(' - predlozuvac', '')
    s = re.sub(r'[^a-zа-яѓѕјљњќџćčšžđћ0-9\s]', ' ', s)
    res = [CYR_TO_LAT.get(ch, ch) for ch in s]
    norm = "".join(res)
    norm = norm.replace('sh', 's').replace('ch', 'c').replace('zh', 'z').replace('dj', 'gj').replace('kj', 'k').replace('ikj', 'ik').replace('ic', 'ik')
    return re.sub(r'\s+', ' ', norm).strip()

ALIAS_MAP = {
    'милка јовевска': ['mila mila', 'milka jovevska'],
    'milka jovevska': ['mila mila', 'milka jovevska'],
    'angela dr mit': ['ani mit', 'angela josifoska'],
    'орхидеја јосифоска': ['ani mit', 'angela josifoska'],
    'restylecloset6': ['да одгледуваме повеќе пченка'],
    'да одгледуваме повеќе пченка': ['да одгледуваме повеќе пченка'],
    'zzz boskovic': ['vojka boshkovic'],
    'vojka boshkovic': ['vojka boshkovic'],
    'daniela_kosta': ['jelica damjanovska'],
    'jelica damjanovska': ['jelica damjanovska'],
    'jovann_d': ['jovan damjanovski'],
    'профил 2': ['jovan damjanovski'],
    'александар петров': ['александар бацка', 'aleksandar backa'],
    'ljupco nikolovski': ['бу бу'],
    'andrijan andrijan': ['andrijan ani'],
    'елена божиновска': ['elena bozhinoska', 'elena bozinoska'],
    'elena bozinovska': ['elena bozhinoska', 'elena bozinoska'],
    'kate arsova velkova': ['kate arsova', 'kate arsova velkova', 'кате арсова'],
    'kate arsova': ['kate arsova', 'kate arsova velkova', 'кате арсова'],
    'емилија колева илиевска': ['koleva emilija'],
    'ѓокески филип': ['filip gjokeski'],
    'slobodan damkovski': ['bobi damkovski'],
    'рајна ѓорѓиева': ['rajne rajni'],
    'rajna gjorgjieva': ['rajne rajni'],
    'игор милошевски': ['м.игор ххх', 'м.игор xxx'],
    'igor milosevski': ['м.игор ххх', 'м.игор xxx'],
    'ацка станишева': ['ацка станишева', 'elena stanisheva'],
    'elena stanisheva': ['ацка станишева'],
    'ceki s medenica': ['ceki s medenica-dimeska'],
    'орце симов': ['orce simov nov profil'],
    'orce simov': ['orce simov nov profil'],
    'марија талевска': ['marija talevska miche'],
    'marija talevska': ['marija talevska miche'],
    'драган спировски': ['dragan spirovski gagi'],
    'драган димитров': ['драган гаги димитров'],
    'софија вренцоска': ['софи вренцоска'],
    'sofi vrencoska': ['софи вренцоска'],
    'владимир спиров': ['vladimir spirov'],
    'маја стоиманова': ['maja stoimanova-ristovska'],
    'снежана златковска': ['snezana zlatkovska'],
    'андријана филиповска': ['andrijana filiposka'],
    'andrijana filipovska': ['andrijana filiposka'],
    'jadranka popovska': ['jadranka popovski'],
    'јадранка поповска': ['jadranka popovski'],
    'katerina avramovska markovik': ['katerina sotirova', 'sotirova-nikolova katerina', 'kate petrovska-markovic'],
    'катерина аврамовска марковиќ': ['kate petrovska-markovic', 'sotirova-nikolova katerina'],
    'катерина сотирова': ['sotirova-nikolova katerina'],
    'nikolinka aneva': ['nina aneva'],
    'весна вучковска': ['vuckovska vesna'],
    'петреска јулијана': ['petrovska julijana'],
    'јулијана петровска': ['petrovska julijana'],
    'стојан срезоски': ['srezoski stojan'],
    'оливера костадиновска': ['olivera kostadinoska'],
    'димитри таки талевски': ['dimitri taki talevski'],
    'марија петкоски': ['marija petkovska'],
    'гоце талевски': ['goce talevski'],
    'дарко димитров ерасмус': ['darko dimitrov'],
    'darko dimitrov erasmus': ['darko dimitrov'],
    'илина шандуловска': ['ilina shandulovska'],
    'билијана блажевска': ['блажевска билјана'],
    'билјана блажевска': ['блажевска билјана'],
    'zokica_b': ['zorana boskovic'],
    'zorana': ['zorana boskovic'],
    'dimcedamjanovski': ['dimce damjanovski'],
    'зага костадиновска': ['dragica kostadinovska', 'зага костадиновска'],
    'daniel cvetanoski': ['daniel cvet'],
    'dragojuv miljkovikj': ['miljkovic dragoljub'],
    'jasna jasna': ['jas jas'],
    'olgica lazareska': ['lazarevska olga'],
    'еко планета': ['еко планета', 'ekoplaneta'],
    'кристијан ј.': ['кристијан ј.'],
    'kristijan j.': ['кристијан ј.'],
    'би лј ана': ['би лј ана'],
    'maja ristevska ipard': ['maja ristevska'],
    'sonja dimovska': ['sonja di', 'sanja dimovska'],
    'стефан жути': ['стефан жути'],
    'aco mickoski': ['ace mickoski'],
    'beti zhivadinovik': ['beti b. zivadinovik'],
    'олга образование': ['olga dukovska-atov', 'olga dukovska atov'],
    'мирјана нанчева дуковска': ['mirjana nanceva-dukovska', 'olga dukovska-atov']
}

NAME_MAP_CYR = {
    "Boro": "Боро", "Predrag": "Предраг", "Andrijan": "Андријан", 
    "Darko": "Дарко", "Daniela": "Даниела", "Beti": "Бети", 
    "Nadica": "Надица", "Jadranka": "Јадранка", "Maja": "Маја", 
    "Biljana": "Билјана", "Anita": "Анита", "Vlatce": "Влатче",
    "Goran": "Горан", "Dejan": "Дејан", "Emilija": "Емилија",
    "Katerina": "Катерина", "Kristina": "Кристина", "Olga": "Олга",
    "Olgica": "Олгица", "Rajna": "Рајна", "Cvetanka": "Цветанка",
    "Andrej": "Андреј", "Mihail": "Михаил", "Elena": "Елена",
    "Spiro": "Спиро", "Mitko": "Митко", "Vlatko": "Влатко",
    "Lara": "Лара", "Damjan": "Дамјан", "Petar": "Петар",
    "Martina": "Мартина", "Slavica": "Славица", "Dali": "Дали",
    "Gorgi": "Ѓорѓи", "Dimitar": "Димитар", "Teona": "Теона"
}

LINKS_META_FILE = "post_links_meta.json"

def load_links_meta():
    default_meta = {
        "current_link": "https://www.facebook.com/share/18VUNHn6g6/",
        "current_title": "Актуелна ФБ Објава (Провери пред праќање)",
        "history": [
            {"post_col": "Дваесет и втора Објава (803 Конечен Пресек)", "url": "https://www.facebook.com/share/18VUNHn6g6/", "time": "2026-10-21"},
            {"post_col": "Дваесет и трета Објава (463 Пресек)", "url": "https://www.facebook.com/share/18VUNHn6g6/", "time": "2026-10-22"}
        ]
    }
    if os.path.exists(LINKS_META_FILE):
        try:
            with open(LINKS_META_FILE, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            return default_meta
    return default_meta

def save_links_meta(meta_data):
    try:
        with open(LINKS_META_FILE, "w", encoding="utf-8") as f:
            json.dump(meta_data, f, ensure_ascii=False, indent=2)
    except Exception:
        pass

# ==============================================================================
# 2. ФУНКЦИИ ЗА НАОЃАЊЕ И КОМПЛЕТНО ГЕНЕРИРАЊЕ НА EXCEL БАЗАТА (7 ЛИСТОВИ)
# ==============================================================================

# ==============================================================================
# GITHUB AUTO-SYNC ФУНКЦИЈА
# ==============================================================================
def push_to_github(file_path):
    import base64
    import json
    import urllib.request
    import urllib.error

    try:
        token = st.secrets.get("GITHUB_TOKEN")
        repo = st.secrets.get("GITHUB_REPO")
        if not token or not repo:
            return False, "Не се пронајдени GitHub Secrets."
    except Exception:
        return False, "Секретите не се конфигурирани."

    try:
        with open(file_path, "rb") as f:
            content = f.read()
        encoded_content = base64.b64encode(content).decode("utf-8")
        
        url = f"https://api.github.com/repos/{repo}/contents/{file_path}"
        headers = {
            "Authorization": f"Bearer {token}",
            "Accept": "application/vnd.github.v3+json",
            "User-Agent": "Streamlit-Auto-Sync"
        }
        
        # Check if file exists to get its SHA
        req = urllib.request.Request(url, headers=headers)
        sha = None
        try:
            with urllib.request.urlopen(req) as response:
                if response.status == 200:
                    resp_data = json.loads(response.read().decode())
                    sha = resp_data.get("sha")
        except urllib.error.HTTPError as e:
            if e.code != 404:
                return False, f"Грешка при проверка на фајлот: {e.code}"
                
        data = {
            "message": f"Автоматска синхронизација на {file_path} преку Streamlit",
            "content": encoded_content
        }
        if sha:
            data["sha"] = sha
            
        req_put = urllib.request.Request(url, data=json.dumps(data).encode("utf-8"), headers=headers, method="PUT")
        try:
            with urllib.request.urlopen(req_put) as response:
                if response.status in [200, 201]:
                    return True, "Базата е успешно и автоматски зачувана на GitHub!"
        except urllib.error.HTTPError as e:
            return False, f"Грешка при синхронизација: {e.code} - {e.read().decode()}"
            
    except Exception as e:
        return False, f"Проблем: {str(e)}"
    
    return False, "Непозната грешка при GitHub синхронизација."

def get_latest_excel():
    files = [f for f in os.listdir() if f.endswith('.xlsx') and 'kontrola' in f and not f.startswith('~$')]
    if not files:
        return None
    def get_ver(fn):
        nums = re.findall(r'\d+', fn)
        return int(nums[-1]) if nums else 0
    files.sort(key=get_ver)
    return files[-1]

def get_next_version_filename(current_file):
    if not current_file:
        return "kontrola_klik_najnova_final_219.xlsx"
    nums = re.findall(r'\d+', current_file)
    next_ver = (int(nums[-1]) + 1) if nums else 219
    return f"kontrola_klik_najnova_final_{next_ver}.xlsx"

def get_eval_columns(df_main):
    cols = []
    keywords = ['Објава', 'Objava', 'Проверка', 'Пресек']
    start_tracking = False
    for c in df_main.columns:
        if 'Дваесет и втора Објава' in c:
            start_tracking = True
        if any(k in c for k in keywords) or start_tracking:
            if c not in ['Р.б.', 'Име и презиме / Профил', 'Предложувач', 'Име на Телеграм (За пребарување)', 'Телефонски број', 'Одговорен контролор', 'Телефонски број (Telegram)', 'Тековен Пресек']:
                cols.append(c)
    return cols

def get_official_posts(df_main):
    eval_cols = get_eval_columns(df_main)
    base_official = [
        'Проверка 23 (306 реакции)',
        'Нова Објава 2 (761 реакции)',
        'Трета Објава 4 (320 реакции)',
        'Четврта Објава 3 (771 реакции)',
        'Петта Објава 2 (637 реакции)',
        'Шеста Објава (Ревидиран Пресек)',
        'Седма Објава (Длабински Фузи Пресек)',
        'Осма Објава (Ново Видео)',
        'Десетта Објава (Најнов CSV - Финална Проверка)',
        'Единаесетта Објава (FAST CSV)',
        'Дванаесетта Објава (Нов Пресек)',
        'Тринаесетта Објава (Видео Пресек)',
        'Четринаесетта Објава (Реално Видео)',
        'Петринаесетта Објава (Ново Видео)',
        'Седумнаесетта Објава (Конечен Пресек)',
        'Осумнаесетта Објава (Конечен Пресек)',
        'Деветнаесетта Објава (1.1k Пресек)',
        'Дваесетта Објава (919 Пресек)',
        'Дваесет и прва Објава (1K+ Пресек - По Потсетници)',
        'Дваесет и втора Објава (803 Конечен Пресек)'
    ]
    valid_cols = [c for c in base_official if c in df_main.columns]
    # Додај ги сите понови објави по 22-рата
    if 'Дваесет и втора Објава (803 Конечен Пресек)' in eval_cols:
        idx_22 = eval_cols.index('Дваесет и втора Објава (803 Конечен Пресек)')
        for c in eval_cols[idx_22 + 1:]:
            if c not in valid_cols:
                valid_cols.append(c)
    else:
        for c in eval_cols:
            if c not in valid_cols:
                valid_cols.append(c)
    return valid_cols

def build_all_tables(df_main):
    official_cols = get_official_posts(df_main)
    eval_cols = get_eval_columns(df_main)
    col_cur = eval_cols[-1]
    col_prev = eval_cols[-2] if len(eval_cols) >= 2 else col_cur

    stats_list = []
    for proposer, grp in df_main.groupby('Предложувач'):
        tot_members = len(grp)
        p_clean = str(proposer).split(' (http')[0].replace('[', '').replace(']', '').replace(' - Предложувач', '').replace(' - Predlozuvac', '').strip()
        tg_n = grp['Име на Телеграм (За пребарување)'].iloc[0] if 'Име на Телеграм (За пребарување)' in grp.columns else ''

        post_likes = []
        post_effs = []
        for col in official_cols:
            l = (grp[col].astype(str).str.contains('✅', na=False)).sum()
            post_likes.append(l)
            post_effs.append(l / tot_members if tot_members > 0 else 0)

        first_act_idx = next((i for i, l in enumerate(post_likes) if l > 0), None)
        if first_act_idx is not None:
            rel_likes = post_likes[first_act_idx:]
            rel_effs = post_effs[first_act_idx:]
        else:
            rel_likes = [post_likes[-1]] if post_likes else [0]
            rel_effs = [post_effs[-1]] if post_effs else [0]

        avg_likes = sum(rel_likes) / len(rel_likes)
        avg_eff = sum(rel_effs) / len(rel_effs)

        prev_likes = (grp[col_prev].astype(str).str.contains('✅', na=False)).sum()
        cur_likes = (grp[col_cur].astype(str).str.contains('✅', na=False)).sum()
        diff_likes = cur_likes - prev_likes
        cur_eff = cur_likes / tot_members if tot_members > 0 else 0
        cur_miss = tot_members - cur_likes

        eff_pct = avg_eff * 100
        if eff_pct <= 30.0:
            gr_hist = 'I'
        elif eff_pct < 50.0:
            gr_hist = 'II'
        elif eff_pct < 65.0:
            gr_hist = 'III'
        else:
            gr_hist = 'IV'

        stats_list.append({
            'Предлагач': p_clean,
            'ГР.': gr_hist,
            'Телеграм назив': tg_n,
            'Квантум': tot_members,
            'Претходна': prev_likes,
            'Тековна': cur_likes,
            'Разлика (+/-)': diff_likes,
            'Тековни Фалат': cur_miss,
            'Тековна Ефикасност': cur_eff,
            'Тековна Ефикасност %': f"{cur_eff*100:.1f}%",
            'Просек Лајкови': round(avg_likes, 1),
            'Просечна Ефикасност': avg_eff,
            'Просечна Ефикасност %': f"{avg_eff*100:.1f}%",
            'Објави Следени': len(rel_likes)
        })

    df_comp = pd.DataFrame(stats_list)

    # 1. Sumarno
    df_sum_disp = df_comp.sort_values(
        by=['Тековна', 'Тековна Ефикасност', 'Квантум'],
        ascending=[False, False, False]
    ).reset_index(drop=True)
    df_sum_disp.insert(0, 'Ранг', range(1, len(df_sum_disp) + 1))

    df_sum_sheet = pd.DataFrame({
        'Ранг': df_sum_disp['Ранг'],
        'ГР.': df_sum_disp['ГР.'],
        'Предлагач': df_sum_disp['Предлагач'],
        'Претходна': df_sum_disp['Претходна'],
        'Тековна': df_sum_disp['Тековна'],
        'Разлика (+/-)': df_sum_disp['Разлика (+/-)'],
        'Фалат': df_sum_disp['Тековни Фалат'],
        'Квантум': df_sum_disp['Квантум'],
        'Ефикасност %': df_sum_disp['Тековна Ефикасност %'],
        'Телеграм назив': df_sum_disp['Телеграм назив']
    })

    # 2. Групи IV - I
    def make_group_df(gr_code):
        sub = df_comp[df_comp['ГР.'] == gr_code].copy()
        sub = sub.sort_values(by=['Просек Лајкови', 'Просечна Ефикасност', 'Квантум'], ascending=[False, False, False]).reset_index(drop=True)
        sub.insert(0, 'Поз.', range(1, len(sub) + 1))
        status_col = []
        for _, r in sub.iterrows():
            cur_v = float(r['Тековна Ефикасност %'].replace('%', ''))
            avg_v = float(r['Просечна Ефикасност %'].replace('%', ''))
            diff = cur_v - avg_v
            if diff >= 10.0:
                status_col.append(f"🟢 Натфрлил (+{diff:.1f}%)")
            elif diff <= -10.0:
                status_col.append(f"🔴 Потфрлил ({diff:.1f}%)")
            else:
                status_col.append(f"⚪ Стандардно ({diff:+.1f}%)")
        return pd.DataFrame({
            'Поз.': sub['Поз.'],
            'Предлагач': sub['Предлагач'],
            'Квантум': sub['Квантум'],
            'Просек Лајкови': sub['Просек Лајкови'],
            'Просечна Ефикасност %': sub['Просечна Ефикасност %'],
            'Тековни Активни': sub['Тековна'],
            'Тековни Фалат': sub['Тековни Фалат'],
            'Тековна Ефикасност %': sub['Тековна Ефикасност %'],
            'Статус / Отстапување': status_col,
            'Телеграм назив': sub['Телеграм назив']
        })

    df_g4 = make_group_df('IV')
    df_g3 = make_group_df('III')
    df_g2 = make_group_df('II')
    df_g1 = make_group_df('I')

    # 3. Неактивни (0 Лајкови на последните 5 објави)
    last_5_official = official_cols[-5:] if len(official_cols) >= 5 else official_cols
    inactive_rows = []
    for _, r in df_main.iterrows():
        likes_all = sum(1 for c in eval_cols if '✅' in str(r[c]))
        likes_hist = sum(1 for c in official_cols if '✅' in str(r[c]))
        likes_5 = sum(1 for c in last_5_official if '✅' in str(r[c]))
        if likes_5 == 0:
            status_inact = "🔴 0 лајкови од почеток (Кандидат за замена/проверка)" if likes_all == 0 else f"🟠 Неактивен на последните 5 објави (Претходно: {likes_hist} лајкови)"
            clean_p = str(r['Предложувач']).split(' (http')[0].replace('[', '').replace(']', '').replace(' - Предложувач', '').replace(' - Predlozuvac', '').strip()
            clean_m = str(r['Име и презиме / Профил']).split(' (http')[0].strip()
            inactive_rows.append({
                'Р.б. во База': r['Р.б.'],
                'Предлагач': clean_p,
                'Име и презиме / Профил (Член)': clean_m,
                'Лајкови (Последни 5 објави)': likes_5,
                'Лајкови (Вкупно историски)': likes_hist,
                'Дијагноза / Препорака': status_inact,
                'Телеграм назив на предлагач': r['Име на Телеграм (За пребарување)']
            })
    df_inactive = pd.DataFrame(inactive_rows)
    if not df_inactive.empty:
        df_inactive = df_inactive.sort_values(by=['Лајкови (Вкупно историски)', 'Предлагач', 'Р.б. во База'], ascending=[True, True, True]).reset_index(drop=True)
        df_inactive.insert(0, 'Ред.', range(1, len(df_inactive) + 1))

    return {
        'Glavna Baza': df_main,
        'Sumarno': df_sum_sheet,
        'Grupa IV': df_g4,
        'Grupa III': df_g3,
        'Grupa II': df_g2,
        'Grupa I': df_g1,
        'Neaktivni (0 Lajkovi)': df_inactive,
        '_comp': df_comp
    }

def save_formatted_excel(df_main, output_filename):
    tables = build_all_tables(df_main)
    df_comp = tables['_comp']
    with pd.ExcelWriter(output_filename, engine='openpyxl') as writer:
        tables['Glavna Baza'].to_excel(writer, sheet_name='Glavna Baza', index=False)
        tables['Sumarno'].to_excel(writer, sheet_name='Sumarno', index=False)
        tables['Grupa IV'].to_excel(writer, sheet_name='Grupa IV', index=False)
        tables['Grupa III'].to_excel(writer, sheet_name='Grupa III', index=False)
        tables['Grupa II'].to_excel(writer, sheet_name='Grupa II', index=False)
        tables['Grupa I'].to_excel(writer, sheet_name='Grupa I', index=False)
        tables['Neaktivni (0 Lajkovi)'].to_excel(writer, sheet_name='Neaktivni (0 Lajkovi)', index=False)

    wb = openpyxl.load_workbook(output_filename)
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9')
    )
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    total_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    total_font = Font(name="Calibri", size=11, bold=True, color="000000")
    total_border = Border(
        left=Side(style='thin', color='A6A6A6'), right=Side(style='thin', color='A6A6A6'),
        top=Side(style='thin', color='000000'), bottom=Side(style='double', color='000000')
    )
    right_align_bold = Alignment(horizontal="right", vertical="center")

    group_fills = {
        'Grupa IV': PatternFill(start_color="274E13", end_color="274E13", fill_type="solid"),
        'Grupa III': PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid"),
        'Grupa II': PatternFill(start_color="B45F06", end_color="B45F06", fill_type="solid"),
        'Grupa I': PatternFill(start_color="990000", end_color="990000", fill_type="solid"),
        'Neaktivni (0 Lajkovi)': PatternFill(start_color="7F0000", end_color="7F0000", fill_type="solid")
    }

    # Glavna Baza
    ws_main = wb['Glavna Baza']
    ws_main.freeze_panes = 'A2'
    eval_cols = get_eval_columns(df_main)
    cols_to_show = eval_cols[-3:]
    for col in ws_main.iter_cols(min_row=1, max_row=1, max_col=ws_main.max_column):
        cell = col[0]
        val = str(cell.value)
        if any(k in val for k in ['Proverka', 'Objava', 'Проверка', 'Објава']) and val not in cols_to_show:
            ws_main.column_dimensions[cell.column_letter].hidden = True
    for cell in ws_main[1]:
        cell.fill = header_fill; cell.font = header_font; cell.alignment = center_align
    ws_main.column_dimensions['A'].width = 7
    ws_main.column_dimensions['B'].width = 38
    ws_main.column_dimensions['C'].width = 32

    # Sumarno
    ws_sum = wb['Sumarno']
    ws_sum.freeze_panes = 'A2'
    for cell in ws_sum[1]:
        cell.fill = header_fill; cell.font = header_font; cell.alignment = center_align
    for row in ws_sum.iter_rows(min_row=2, max_row=ws_sum.max_row, max_col=ws_sum.max_column):
        for cell in row:
            cell.border = thin_border
            cell.alignment = center_align if cell.column_letter in ['A', 'B', 'D', 'E', 'F', 'G', 'H', 'I'] else left_align
    tot_cur = df_comp['Тековна'].sum()
    tot_q = df_comp['Квантум'].sum()
    max_r = ws_sum.max_row; tot_r = max_r + 1
    ws_sum.cell(row=tot_r, column=3, value="ВКУПНО:")
    ws_sum.cell(row=tot_r, column=4, value=f"=SUM(D2:D{max_r})")
    ws_sum.cell(row=tot_r, column=5, value=f"=SUM(E2:E{max_r})")
    ws_sum.cell(row=tot_r, column=6, value=f"=SUM(F2:F{max_r})")
    ws_sum.cell(row=tot_r, column=7, value=f"=SUM(G2:G{max_r})")
    ws_sum.cell(row=tot_r, column=8, value=f"=SUM(H2:H{max_r})")
    ws_sum.cell(row=tot_r, column=9, value=f"{tot_cur/tot_q*100:.1f}%" if tot_q > 0 else "0.0%")
    for c in range(1, 11):
        cell = ws_sum.cell(row=tot_r, column=c)
        cell.fill = total_fill; cell.font = total_font; cell.border = total_border
        cell.alignment = right_align_bold if c == 3 else center_align

    # Grupa IV - I
    for sh_name in ['Grupa IV', 'Grupa III', 'Grupa II', 'Grupa I']:
        ws = wb[sh_name]
        ws.freeze_panes = 'A2'
        df_g = tables[sh_name]
        s_fill = group_fills[sh_name]
        for cell in ws[1]:
            cell.fill = s_fill; cell.font = header_font; cell.alignment = center_align
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=ws.max_column):
            for cell in row:
                cell.border = thin_border
                cell.alignment = center_align if cell.column_letter in ['A', 'C', 'D', 'E', 'F', 'G', 'H', 'I'] else left_align
        max_r = ws.max_row; tot_r = max_r + 1
        q_s = df_g['Квантум'].sum(); avg_s = df_g['Просек Лајкови'].sum(); act_s = df_g['Тековни Активни'].sum()
        ws.cell(row=tot_r, column=2, value="ВКУПНО:")
        ws.cell(row=tot_r, column=3, value=f"=SUM(C2:C{max_r})")
        ws.cell(row=tot_r, column=4, value=f"=SUM(D2:D{max_r})")
        ws.cell(row=tot_r, column=5, value=f"{avg_s/q_s*100:.1f}%" if q_s > 0 else "0.0%")
        ws.cell(row=tot_r, column=6, value=f"=SUM(F2:F{max_r})")
        ws.cell(row=tot_r, column=7, value=f"=SUM(G2:G{max_r})")
        ws.cell(row=tot_r, column=8, value=f"{act_s/q_s*100:.1f}%" if q_s > 0 else "0.0%")
        for c in range(1, 11):
            cell = ws.cell(row=tot_r, column=c)
            cell.fill = total_fill; cell.font = total_font; cell.border = total_border
            cell.alignment = right_align_bold if c == 2 else center_align

    # Neaktivni
    ws_inact = wb['Neaktivni (0 Lajkovi)']
    ws_inact.freeze_panes = 'A2'
    for cell in ws_inact[1]:
        cell.fill = group_fills['Neaktivni (0 Lajkovi)']; cell.font = header_font; cell.alignment = center_align

    wb.save(output_filename)
    
    # Автоматска синхронизација со GitHub
    success, msg = push_to_github(output_filename)
    if success:
        st.toast("✅ " + msg)
    else:
        st.toast("⚠️ " + msg)
        
    return tables

# ==============================================================================
# 3. КОМПОНЕНТА ЗА ПРЕГЛЕД И ВЕРИФИКАЦИЈА НА ФБ ЛИНК ОД ОБЈАВА (PREVIEW)
# ==============================================================================
def render_post_link_preview_box(links_meta, eval_cols, key_suffix="main"):
    cur_link = links_meta.get("current_link", "https://www.facebook.com/share/18VUNHn6g6/")
    cur_title = links_meta.get("current_title", "Актуелна ФБ Објава")
    latest_col = eval_cols[-1] if eval_cols else "Нема колона"

    st.markdown(f"""
    <div class="link-card">
        <div style="display:flex; justify-content:space-between; align-items:center; flex-wrap:wrap; gap:8px;">
            <div>
                <span class="link-badge">🔗 АКТИВЕН ЛИНК НА ОБЈАВАТА</span>
                <span style="margin-left:8px; font-weight:bold; color:#1F497D;">{cur_title}</span>
            </div>
            <div style="font-size:13px; color:#333;">
                📌 Поврзана колона во база: <b>{latest_col}</b>
            </div>
        </div>
        <div style="margin-top:10px; font-size:15px; word-break:break-all;">
            👉 Директен линк: <a href="{cur_link}" target="_blank" style="font-weight:bold; color:#0b57d0;">{cur_link}</a>
        </div>
    </div>
    """, unsafe_allow_html=True)

    with st.expander("👁️ Прикажи Визуелен Преглед (Live Preview) и Промени Линк на Објавата", expanded=False):
        c1, c2 = st.columns([3, 2])
        with c1:
            new_url = st.text_input("Внеси / Корегирај линк до ФБ објавата:", value=cur_link, key=f"url_in_{key_suffix}")
            new_title = st.text_input("Кратко име / Опис на објавата (за твоја евиденција):", value=cur_title, key=f"title_in_{key_suffix}")
            if st.button("💾 Зачувај го новиот линк како активен", key=f"save_link_{key_suffix}"):
                links_meta["current_link"] = new_url.strip()
                links_meta["current_title"] = new_title.strip()
                links_meta.setdefault("history", []).append({
                    "post_col": latest_col,
                    "url": new_url.strip(),
                    "title": new_title.strip(),
                    "time": datetime.now().strftime("%Y-%m-%d %H:%M")
                })
                save_links_meta(links_meta)
                st.success("✅ Линкот е успешно зачуван и поставен во генераторот за Телеграм пораки!")
                st.rerun()

            st.markdown(f"🔍 **[Кликни тука за да ја отвориш објавата директно на Facebook во нов прозорец]({new_url})**")
            if links_meta.get("history"):
                st.markdown("**📜 Историја на користени линкови по објави:**")
                st.dataframe(pd.DataFrame(links_meta["history"][::-1]), use_container_width=True, height=150)

        with c2:
            st.markdown("**🖼️ Слика од објавата (Screenshot Preview):**")
            st.caption("Прикачи скриншот од објавата за полесна визуелна контрола.")
            
            uploaded_img = st.file_uploader("Upload screenshot:", type=['png', 'jpg', 'jpeg'], key=f"img_up_{key_suffix}")
            preview_file = "active_post_preview.png"
            
            if uploaded_img is not None:
                with open(preview_file, "wb") as f:
                    f.write(uploaded_img.getbuffer())
                st.success("Скриншотот е успешно зачуван!")
                st.rerun()
                
            if os.path.exists(preview_file):
                st.image(preview_file, caption="Тековна објава", use_container_width=True)

# ==============================================================================
# 4. ГЛАВЕН ИНТЕРФЕЈС И ВЧИТУВАЊЕ НА БАЗАТА
# ==============================================================================
st.markdown('<div class="main-header">📊 ФБ Контрола и Менаџмент на Мрежата (Десктоп & Мобилен)</div>', unsafe_allow_html=True)

links_meta = load_links_meta()
latest_file = get_latest_excel()

# Sidebar контрола на фајлот
st.sidebar.markdown("### 🗄️ Мастер База")
uploaded_master = st.sidebar.file_uploader("Прикачи понова база (.xlsx) доколку сакаш:", type=['xlsx'])
if uploaded_master is not None:
    if "last_uploaded_file" not in st.session_state or st.session_state["last_uploaded_file"] != uploaded_master.name:
        save_name = uploaded_master.name if 'kontrola' in uploaded_master.name else get_next_version_filename(latest_file)
        with open(save_name, "wb") as f:
            f.write(uploaded_master.getbuffer())
        st.session_state["last_uploaded_file"] = uploaded_master.name
        latest_file = save_name
        st.sidebar.success(f"✅ Вчитана нова база: {latest_file}")
    else:
        # File was already uploaded and processed; let the system detect newer versions (like 221, 222) naturally
        latest_file = get_latest_excel()

if not latest_file:
    st.error("❌ Не е пронајден Excel фајл со мастер базата (`kontrola_klik_...xlsx`)! Прикачи го во левото мени.")
    st.stop()

st.sidebar.success(f"📂 Активна верзија: **{latest_file}**")
# Копче за симнување на најновата база во секое време
with open(latest_file, "rb") as f_dl:
    st.sidebar.download_button(
        label="📥 Преземи ја Мастер Базата (.xlsx)",
        data=f_dl,
        file_name=latest_file,
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True
    )

xls_active = pd.ExcelFile(latest_file)
df_main = pd.read_excel(xls_active, sheet_name='Glavna Baza')
tables_data = build_all_tables(df_main)
eval_cols = get_eval_columns(df_main)

# Главно Мени
menu = st.sidebar.radio(
    "📌 Навигација:",
    [
        "📊 1. Контролна Табла & Сите Листови",
        "⚡ 2. Фузи Пресек (Нова Објава / Потсетник)",
        "👥 3. Управување со Тимови (Додај / Избриши)",
        "📨 4. Телеграм Потсетници (Две Брзини)",
        "🔍 5. Пребарување & Алијас Речник",
        "📱 6. Мобилен Пристап & Икона"
    ]
)

# ------------------------------------------------------------------------------
# ТАБ 1: КОНТРОЛНА ТАБЛА И СИТЕ 7 ЛИСТОВИ (SHEETS)
# ------------------------------------------------------------------------------
if menu.startswith("📊 1."):
    render_post_link_preview_box(links_meta, eval_cols, key_suffix="tab1")

    total_members = len(df_main)
    df_comp = tables_data['_comp']
    total_teams = len(df_comp)
    latest_col = eval_cols[-1]
    prev_col = eval_cols[-2] if len(eval_cols) >= 2 else latest_col

    active_cur = int((df_main[latest_col].astype(str).str.contains('✅', na=False)).sum())
    active_prev = int((df_main[prev_col].astype(str).str.contains('✅', na=False)).sum())
    diff_act = active_cur - active_prev
    eff_rate = (active_cur / total_members * 100) if total_members > 0 else 0

    m1, m2, m3, m4 = st.columns(4)
    m1.metric("👥 Вкупен Квантум (Членови)", f"{total_members}")
    m2.metric("🏆 Активни Тимови", f"{total_teams}")
    m3.metric(f"✅ Активни ({latest_col[:22]}...)", f"{active_cur}", delta=f"{diff_act:+d} од претходен пресек")
    m4.metric("📈 Тековна Ефикасност", f"{eff_rate:.1f}%", delta=f"Фалат: {total_members - active_cur}")

    st.markdown("---")
    sheet_tabs = st.tabs([
        "📈 Sumarno (Ранг Листа)",
        "🟢 Grupa IV (≥65%)",
        "🔵 Grupa III (50-65%)",
        "🟠 Grupa II (30-50%)",
        "🔴 Grupa I (≤30%)",
        "⚠️ Neaktivni (0 Лајкови)",
        "🗄️ Glavna Baza (Сите)"
    ])

    with sheet_tabs[0]:
        st.markdown(f"**📋 Сумарна Ранг Листа — Споредба: `{prev_col}` ➡️ `{latest_col}`**")
        st.dataframe(tables_data['Sumarno'], use_container_width=True, height=520)
        st.info(f"**ВКУПНО МРЕЖА:** Претходна: **{active_prev}** | Тековна: **{active_cur}** | Разлика: **{diff_act:+d}** | Фалат: **{total_members - active_cur}** | Квантум: **{total_members}** | Ефикасност: **{eff_rate:.1f}%**")

    with sheet_tabs[1]:
        g4 = tables_data['Grupa IV']
        st.markdown(f"**🟢 Група IV — Елитни Тимови ({len(g4)} тимови, квантум: {g4['Квантум'].sum()} члена)**")
        st.dataframe(g4, use_container_width=True, height=480)

    with sheet_tabs[2]:
        g3 = tables_data['Grupa III']
        st.markdown(f"**🔵 Група III — Стабилни Тимови ({len(g3)} тимови, квантум: {g3['Квантум'].sum()} члена)**")
        st.dataframe(g3, use_container_width=True, height=480)

    with sheet_tabs[3]:
        g2 = tables_data['Grupa II']
        st.markdown(f"**🟠 Група II — Развојни Тимови ({len(g2)} тимови, квантум: {g2['Квантум'].sum()} члена)**")
        st.dataframe(g2, use_container_width=True, height=480)

    with sheet_tabs[4]:
        g1 = tables_data['Grupa I']
        st.markdown(f"**🔴 Група I — Критични Тимови ({len(g1)} тимови, квантум: {g1['Квантум'].sum()} члена)**")
        st.dataframe(g1, use_container_width=True, height=480)

    with sheet_tabs[5]:
        inact = tables_data['Neaktivni (0 Lajkovi)']
        st.markdown(f"**⚠️ Членови со 0 реакции на последните 5 објави (Вкупно: {len(inact)} членови)**")
        st.dataframe(inact, use_container_width=True, height=480)

    with sheet_tabs[6]:
        st.markdown("**🗄️ Главна База (Прикажани се последните 3 пресеци):**")
        display_cols = ['Р.б.', 'Име и презиме / Профил', 'Предложувач', 'Име на Телеграм (За пребарување)'] + eval_cols[-3:]
        st.dataframe(df_main[display_cols], use_container_width=True, height=520)

# ------------------------------------------------------------------------------
# ТАБ 2: ДЛАБИНСКИ ФУЗИ-ПРЕСЕК (НОВА ОБЈАВА ИЛИ ПО ПОТСЕТНИЦИ)
# ------------------------------------------------------------------------------
elif menu.startswith("⚡ 2."):
    st.markdown('<div class="sub-header">⚡ Автоматски Длабински Фузи-Пресек (Без Ограничувања)</div>', unsafe_allow_html=True)
    render_post_link_preview_box(links_meta, eval_cols, key_suffix="tab2")

    c_mode1, c_mode2 = st.columns(2)
    with c_mode1:
        presek_mode = st.radio(
            "Избери тип на пресек:",
            [
                "🆕 Нова Објава (Создава нова колона во базата)",
                "🔄 Пресек По Потсетници (Ги додава новите лајкови врз последната колона)"
            ]
        )
    with c_mode2:
        default_col_name = "Дваесет и четврта Објава (Нов Пресек)" if presek_mode.startswith("🆕") else f"{eval_cols[-1]} - По Потсетници"
        target_col_name = st.text_input("Назив на колоната за овој пресек:", value=default_col_name)
        post_link_for_presek = st.text_input("Линк од објавата што ја обработуваме:", value=links_meta.get("current_link", ""))

    st.markdown("#### 📥 Внеси ги имињата од објавата (преку фајл или со директно копирање):")
    in_tab1, in_tab2 = st.tabs(["📂 Прикачи Фајл (iminja_od_video.xlsx / .csv / .txt)", "📋 Залепи Список со Имиња (Copy-Paste)"])

    extracted_names = []
    with in_tab1:
        up_file = st.file_uploader("Избери фајл со имиња кои реагирале:", type=['xlsx', 'csv', 'txt'])
        if up_file is not None:
            if up_file.name.endswith('.xlsx'):
                df_v = pd.read_excel(up_file)
                extracted_names = df_v.iloc[:, 0].dropna().astype(str).str.strip().tolist()
            elif up_file.name.endswith('.csv'):
                df_v = pd.read_csv(up_file)
                extracted_names = df_v.iloc[:, 0].dropna().astype(str).str.strip().tolist()
            else:
                txt = up_file.read().decode('utf-8', errors='ignore')
                extracted_names = [l.strip() for l in txt.splitlines() if l.strip()]

    with in_tab2:
        pasted_text = st.text_area("Залепи ги имињата тука (по едно име во секој ред):", height=180)
        if pasted_text.strip():
            lines = [l.strip() for l in pasted_text.splitlines() if l.strip()]
            # Филтрирај системски зборови како Mention / Add friend / mutual friends
            clean_lines = [l for l in lines if not re.match(r'^(\d+\s+mutual\s+friends|mention|add\s+friend|invite|following|confirm|cancel\s+request)$', l.lower())]
            extracted_names.extend(clean_lines)

    if extracted_names:
        st.success(f"✅ Вчитани се вкупно **{len(extracted_names)}** реакции за фузи-споредба!")

        if st.button("🚀 Изврши Длабински Фузи-Пресек и Зачувај Нова Верзија на Базата", type="primary"):
            norm_video = [(vn, normalize_name(vn), set(normalize_name(vn).split())) for vn in extracted_names]
            last_col = eval_cols[-1]
            if presek_mode.startswith("🔄"):
                matched_indices = set(df_main[df_main[last_col] == '✅'].index)
            else:
                matched_indices = set()

            newly_matched = []
            for idx, row in df_main.iterrows():
                if idx in matched_indices:
                    continue
                raw_m = str(row['Име и презиме / Профил'])
                raw_p = str(row['Предложувач'])
                clean_low = raw_m.lower()
                norm_m = normalize_name(raw_m)
                tokens_m = set(norm_m.split())
                matched = False
                matched_vn = None

                # 1. Проверка во Алијас Речник
                for k, v_list in ALIAS_MAP.items():
                    if k in clean_low:
                        for vt in v_list:
                            for vn, n_vn, _ in norm_video:
                                if vt in vn.lower() or normalize_name(vt) in n_vn:
                                    matched = True; matched_vn = vn; break
                            if matched:
                                break
                    if matched:
                        break

                # 2. Специјална проверка за двоен профил кај Angela dr mit
                if not matched and 'angela dr mit' in raw_p.lower() and ('профил 2' in clean_low or 'орхидеја' in clean_low):
                    for vn, _, _ in norm_video:
                        if 'angela josifoska' in vn.lower() or 'ani mit' in vn.lower():
                            matched = True; matched_vn = vn; break

                # 3. Точно или подмножество на токени
                if not matched and norm_m:
                    for vn, n_vn, t_vn in norm_video:
                        if norm_m == n_vn or (len(tokens_m) >= 2 and len(t_vn) >= 2 and (tokens_m == t_vn or tokens_m.issubset(t_vn) or t_vn.issubset(tokens_m))):
                            matched = True; matched_vn = vn; break

                # 4. Фузи совпаѓање со SequenceMatcher >= 0.83
                if not matched and len(tokens_m) >= 2:
                    mlist = list(tokens_m)
                    for vn, n_vn, t_vn in norm_video:
                        vlist = list(t_vn)
                        if len(vlist) >= 2:
                            sim_c = sum(1 for mt in mlist if len(mt) >= 3 and max([SequenceMatcher(None, mt, vt).ratio() for vt in vlist if len(vt) >= 3] + [0]) >= 0.83)
                            if sim_c >= 2:
                                matched = True; matched_vn = vn; break

                if matched:
                    matched_indices.add(idx)
                    newly_matched.append({
                        'Р.б.': row['Р.б.'],
                        'Член во База': raw_m,
                        'Предлагач': raw_p,
                        'Совпаднато со (Видео)': matched_vn
                    })

            if presek_mode.startswith("🔄"):
                df_main.rename(columns={last_col: target_col_name}, inplace=True)
                df_main[target_col_name] = ['✅' if i in matched_indices else '❌' for i in df_main.index]
            else:
                df_main[target_col_name] = ['✅' if i in matched_indices else '❌' for i in df_main.index]

            next_file = get_next_version_filename(latest_file)
            save_formatted_excel(df_main, next_file)

            # Зачувај го и линкот во историјата
            if post_link_for_presek.strip():
                links_meta["current_link"] = post_link_for_presek.strip()
                links_meta.setdefault("history", []).append({
                    "post_col": target_col_name,
                    "url": post_link_for_presek.strip(),
                    "title": f"Пресек ({len(extracted_names)} реакции)",
                    "time": datetime.now().strftime("%Y-%m-%d %H:%M")
                })
                save_links_meta(links_meta)

            st.balloons()
            st.success(f"🎉 Успешно завршен пресек! Зачувана е новата верзија: **{next_file}** со вкупно **{len(matched_indices)}** активни членови (`✅`) од {len(df_main)} ({len(matched_indices)/len(df_main)*100:.1f}%)!")

            with open(next_file, "rb") as f_new:
                st.download_button(
                    label=f"📥 Преземи ја веднаш новата база: {next_file}",
                    data=f_new,
                    file_name=next_file,
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    type="primary"
                )

            if newly_matched:
                st.markdown(f"#### 🔍 Преглед на новооткриените {len(newly_matched)} совпаѓања во овој пресек:")
                st.dataframe(pd.DataFrame(newly_matched), use_container_width=True)

# ------------------------------------------------------------------------------
# ТАБ 3: УПРАВУВАЊЕ СО ТИМОВИ И ЧЛЕНОВИ (ДОДАЈ / ЗАМЕНИ / ИЗБРИШИ)
# ------------------------------------------------------------------------------
elif menu.startswith("👥 3."):
    st.markdown('<div class="sub-header">👥 Интерактивно Управување со Тимови, Предлагачи и Членови</div>', unsafe_allow_html=True)
    st.caption("Секоја промена автоматски ги пресметува сите 7 листови, прави ренумерирање на Р.б. и создава нова верзија на Excel фајлот.")

    crud_tabs = st.tabs([
        "➕ 1. Додај Член во Тим",
        "🆕 2. Додај Нов Предлагач (Тим)",
        "🔄 3. Замени Член / Промени Статус (✅/❌)",
        "🗑️ 4. Избриши Член од Базата"
    ])

    proposers = sorted(df_main['Предложувач'].dropna().unique().tolist())
    latest_col = eval_cols[-1]

    # 1. Додај Член во постоечки тим
    with crud_tabs[0]:
        with st.form("form_add_member"):
            sel_p = st.selectbox("Избери Предлагач (Тим):", proposers)
            new_m_names = st.text_area("Внеси име на нов член (или повеќе членови, по еден во секој ред):")
            already_clicked = st.checkbox(f"Означи ги со ✅ како веќе активни на тековната објава ({latest_col})", value=False)
            btn_add_m = st.form_submit_button("➕ Додај во Тимот и Зачувај Нова База")

            if btn_add_m and new_m_names.strip():
                names_to_add = [n.strip() for n in new_m_names.splitlines() if n.strip()]
                sample = df_main[df_main['Предложувач'] == sel_p].iloc[0]
                new_rows = []
                for nm in names_to_add:
                    r = {c: '❌' for c in df_main.columns}
                    r['Р.б.'] = len(df_main) + len(new_rows) + 1
                    r['Име и презиме / Профил'] = nm
                    r['Предложувач'] = sel_p
                    for meta_c in ['Телефонски број (Telegram)', 'Одговорен контролор', 'Телефонски број', 'Име на Телеграм (За пребарување)']:
                        if meta_c in df_main.columns:
                            r[meta_c] = sample[meta_c]
                    if already_clicked:
                        r[latest_col] = '✅'
                    new_rows.append(r)

                df_updated = pd.concat([df_main, pd.DataFrame(new_rows)], ignore_index=True)
                df_updated['Р.б.'] = range(1, len(df_updated) + 1)
                next_file = get_next_version_filename(latest_file)
                save_formatted_excel(df_updated, next_file)
                st.success(f"✅ Додадени се {len(names_to_add)} членови кај **{sel_p}**! Зачувано во **{next_file}**.")
                st.rerun()

    # 2. Додај Нов Предлагач (Нов Тим)
    with crud_tabs[1]:
        with st.form("form_add_proposer"):
            new_prop_name = st.text_input("Име и презиме на новиот Предлагач (пр. Damjan Blazevski):")
            new_prop_tg = st.text_input("Име на Телеграм (За пребарување):")
            prop_clicks_too = st.checkbox("Предлагачот и самиот клика (додај ред '[Име] - Предложувач')", value=True)
            team_list_raw = st.text_area("Внеси ги членовите на неговиот тим (по еден во секој ред):")
            btn_add_team = st.form_submit_button("🆕 Креирај Нов Тим и Зачувај База")

            if btn_add_team and new_prop_name.strip():
                p_clean = new_prop_name.strip()
                tg_clean = new_prop_tg.strip() if new_prop_tg.strip() else p_clean
                members_list = [m.strip() for m in team_list_raw.splitlines() if m.strip()]
                if prop_clicks_too:
                    members_list.insert(0, f"[{p_clean}] - Предложувач")

                new_rows = []
                for nm in members_list:
                    r = {c: '❌' for c in df_main.columns}
                    r['Р.б.'] = len(df_main) + len(new_rows) + 1
                    r['Име и презиме / Профил'] = nm
                    r['Предложувач'] = p_clean
                    r['Одговорен контролор'] = "Тони Јаревски"
                    r['Име на Телеграм (За пребарување)'] = tg_clean
                    new_rows.append(r)

                df_updated = pd.concat([df_main, pd.DataFrame(new_rows)], ignore_index=True)
                df_updated['Р.б.'] = range(1, len(df_updated) + 1)
                next_file = get_next_version_filename(latest_file)
                save_formatted_excel(df_updated, next_file)
                st.success(f"🎉 Креиран е нов тим на **{p_clean}** со {len(new_rows)} члена во **{next_file}**!")
                st.rerun()

    # 3. Замени Член или Промени Статус (✅/❌)
    with crud_tabs[2]:
        sel_p_edit = st.selectbox("Избери Тим за измена:", proposers, key="edit_prop_sel")
        team_sub = df_main[df_main['Предложувач'] == sel_p_edit]
        st.dataframe(team_sub[['Р.б.', 'Име и презиме / Профил', latest_col]], use_container_width=True)

        member_options = {f"[{r['Р.б.']}] {r['Име и презиме / Профил']} ({r[latest_col]})": idx for idx, r in team_sub.iterrows()}
        chosen_label = st.selectbox("Избери член од тимот:", list(member_options.keys()))
        chosen_idx = member_options[chosen_label]

        with st.form("form_edit_member"):
            current_name_val = str(df_main.loc[chosen_idx, 'Име и презиме / Профил'])
            current_status_val = str(df_main.loc[chosen_idx, latest_col])
            updated_name = st.text_input("Ново име / Профил (остави исто ако менуваш само статус):", value=current_name_val)
            updated_status = st.selectbox(f"Статус за '{latest_col}':", ['✅', '❌'], index=0 if '✅' in current_status_val else 1)
            reset_history = st.checkbox("Ако ова е целосна замена со нов човек, ресетирај ги старите лајкови на ❌", value=False)
            btn_save_edit = st.form_submit_button("💾 Зачувај Измена")

            if btn_save_edit:
                df_updated = df_main.copy()
                if reset_history:
                    for c in eval_cols:
                        df_updated.loc[chosen_idx, c] = '❌'
                df_updated.loc[chosen_idx, 'Име и презиме / Профил'] = updated_name.strip()
                df_updated.loc[chosen_idx, latest_col] = updated_status
                next_file = get_next_version_filename(latest_file)
                save_formatted_excel(df_updated, next_file)
                st.success(f"✅ Успешно ажуриран член во **{next_file}**!")
                st.rerun()

    # 4. Избриши Член од Базата
    with crud_tabs[3]:
        sel_p_del = st.selectbox("Избери Тим од кој се брише:", proposers, key="del_prop_sel")
        team_sub_del = df_main[df_main['Предложувач'] == sel_p_del]
        del_options = {f"[{r['Р.б.']}] {r['Име и презиме / Профил']}": idx for idx, r in team_sub_del.iterrows()}
        chosen_to_delete = st.multiselect("Избери еден или повеќе членови за бришење:", list(del_options.keys()))

        if st.button("🗑️ Избриши ги избраните членови и Зачувај База", type="primary"):
            if chosen_to_delete:
                idxs = [del_options[k] for k in chosen_to_delete]
                df_updated = df_main.drop(index=idxs).reset_index(drop=True)
                df_updated['Р.б.'] = range(1, len(df_updated) + 1)
                next_file = get_next_version_filename(latest_file)
                save_formatted_excel(df_updated, next_file)
                st.success(f"✅ Избришани се {len(idxs)} членови! Зачувана е нова верзија: **{next_file}**.")
                st.rerun()
            else:
                st.warning("Прво избери барем еден член за бришење.")

# ------------------------------------------------------------------------------
# ТАБ 4: ТЕЛЕГРАМ ПОТСЕТНИЦИ ВО ДВЕ БРЗИНИ + ПРЕГЛЕД НА ЛИНК + PRAKAC.PY
# ------------------------------------------------------------------------------
elif menu.startswith("📨 4."):
    st.markdown('<div class="sub-header">📨 Телеграм Потсетници и Масовни Пораки</div>', unsafe_allow_html=True)

    tg_tabs = st.tabs(["🤖 Автоматски Потсетници (За најнова објава)", "✍️ Прати Сопствена Порака (Custom)"])

    with tg_tabs[0]:
        st.info("Задолжителна контрола: Провери го линкот во картичката подолу пред да ги пуштиш пораките!")
        render_post_link_preview_box(links_meta, eval_cols, key_suffix="tab4")

        link_objava = links_meta.get("current_link", "https://www.facebook.com/share/18VUNHn6g6/")
        col_new = eval_cols[-1]
        last_5_cols = eval_cols[-5:] if len(eval_cols) >= 5 else eval_cols

        messages_data = []
        skipped_teams = []

        for proposer, group in df_main.groupby('Предложувач'):
            total = len(group)
            active = len(group[group[col_new] == '✅'])
            missing = total - active
            eff_pct = (active / total * 100) if total > 0 else 0
            tg_name = group['Име на Телеграм (За пребарување)'].iloc[0]

            if missing > 2 and pd.notna(tg_name) and str(tg_name).strip() != '':
                missing_df = group[group[col_new] != '✅']
                missing_formatted = []
                chronic_zero_count = 0

                for _, row_m in missing_df.iterrows():
                    n = row_m['Име и презиме / Профил']
                    cl = str(n).replace('[', '').replace(']', '').replace(' - Предложувач', '').replace(' - Predlozuvac', '').strip()
                    cl = cl.split(' (http')[0].split(' http')[0].strip()
                    likes_last_5 = sum(1 for c in last_5_cols if '✅' in str(row_m[c]))
                    if likes_last_5 == 0:
                        missing_formatted.append(f"  ▫️ {cl} ❗(0 реакции подолго време)")
                        chronic_zero_count += 1
                    else:
                        missing_formatted.append(f"  ▫️ {cl}")

                first_name = str(proposer).split()[0].split('(')[0].strip()
                first_name = NAME_MAP_CYR.get(first_name, first_name)
                miss_str = "\n".join(missing_formatted)
                zero_note = "\n💡 Забелешка: Оние што се означени со ❗ немаат реакција на последните 5 објави — провери дали кликаат од друг профил или кажи со кого да ги замениме.\n" if chronic_zero_count > 0 else ""

                if eff_pct >= 60.0 or missing <= 4:
                    msg = (
                        f"Здраво {first_name}! 👋\n\n"
                        f"🌟 Браво за тимот, одлично газите на најновата објава! 📌\n"
                        f"Моментално сте на високи: ✅ {active} од {total} активни членови ({eff_pct:.0f}%).\n\n"
                        f"🏁 На чекор сме до 100% — пушти им само уште еден брз потсетник на преостанатите {missing}:\n"
                        f"{miss_str}\n{zero_note}\n🔗 Линк до објавата:\n{link_objava}\n\n"
                        f"Да ве заокружиме на максимум! Ти благодарам! 🚀💪"
                    )
                    brzina_label = "🌟 БРЗИНА 1 (Топ тим)"
                elif active == 0:
                    msg = (
                        f"Здраво {first_name}! 👋\n\n"
                        f"🔔 Важен потсетник за најновата објава 📌\n"
                        f"Од твојот тим моментално сè уште сме на 0 од {total} реакции.\n\n"
                        f"⚠️ Те молам кликни ти прв/а и веднаш препрати го линкот до твојата група за да се приклучите:\n"
                        f"{miss_str}\n{zero_note}\n🔗 Линк до објавата:\n{link_objava}\n\n"
                        f"Ајде да го раздвижиме тимот и да фатиме чекор со мрежата! 🙏🚀"
                    )
                    brzina_label = "⚠️ БРЗИНА 2 (Аларм 0%)"
                else:
                    msg = (
                        f"Здраво {first_name}! 👋\n\n"
                        f"🔔 Потсетник за мобилизација на најновата објава 📌\n"
                        f"Од твојот тим досега се вклучени само: ✅ {active} од {total} членови ({eff_pct:.0f}%).\n\n"
                        f"⚠️ Ни фалат реакции од следните {missing} членови за да го кренеме просекот:\n"
                        f"{miss_str}\n{zero_note}\n🔗 Линк до објавата:\n{link_objava}\n\n"
                        f"Те молам контактирај ги за да се вклучат сите и да ја исполниме целта! 🙏💪"
                    )
                    brzina_label = "🔔 БРЗИНА 2 (Мобилизација)"

                messages_data.append({
                    'Предлагач': proposer,
                    'Телеграм': str(tg_name).strip(),
                    'Брзина': brzina_label,
                    'Активни': f"{active}/{total} ({eff_pct:.0f}%)",
                    'Фалат': missing,
                    'Порака': msg
                })
            else:
                skipped_teams.append({'Предлагач': proposer, 'Активни': f"{active}/{total}", 'Фалат': missing})

        b1_c = sum(1 for m in messages_data if 'БРЗИНА 1' in m['Брзина'])
        b2_m = sum(1 for m in messages_data if 'Мобилизација' in m['Брзина'])
        b2_z = sum(1 for m in messages_data if 'Аларм' in m['Брзина'])

        c1, c2, c3, c4 = st.columns(4)
        c1.metric("📨 Вкупно Пораки (>2 фалат)", len(messages_data))
        c2.metric("🌟 Брзина 1 (Топ тимови)", b1_c)
        c3.metric("🔔 Брзина 2 (Мобилизација)", b2_m)
        c4.metric("⚠️ Брзина 2 (Аларм 0%)", b2_z)

        with st.expander(f"🏆 Поштедени тимови со 0, 1 или 2 што фалат ({len(skipped_teams)} тимови)", expanded=False):
            st.dataframe(pd.DataFrame(skipped_teams), use_container_width=True)

        st.markdown("---")
        st.markdown("#### 🤖 Автоматско испраќање преку Скрипта (prakac.py)")
        if messages_data:
            df_msgs = pd.DataFrame(messages_data)
            csv_data = df_msgs.to_csv(index=False, encoding='utf-8-sig')
            st.download_button(
                label="📥 Преземи CSV фајл за автоматската скрипта (prakac.py)",
                data=csv_data,
                file_name=f"telegram_poraki_za_prakanje.csv",
                mime="text/csv",
                type="primary",
                key="btn_dl_auto_csv"
            )
            st.caption("Преземи го фајлот и искористи го со твојата локална скрипта за да ги пратиш сите пораки одеднаш.")
        else:
            st.info("Нема пораки за испраќање во моментов.")
        st.markdown("---")

        st.markdown("#### 📋 Преглед и Копирање на Поединечни Пораки (Удобно за Мобилен и Десктоп):")
        filter_brzina = st.selectbox("Филтрирај по категорија:", ["Сите пораки", "🌟 БРЗИНА 1 (Топ тим)", "🔔 БРЗИНА 2 (Мобилизација)", "⚠️ БРЗИНА 2 (Аларм 0%)"])
        filtered_msgs = [m for m in messages_data if filter_brzina == "Сите пораки" or m['Брзина'] == filter_brzina]

        sel_msg_prop = st.selectbox("Избери предлагач за преглед и копирање на неговата порака:", [f"{m['Предлагач']} ({m['Активни']} - {m['Брзина']})" for m in filtered_msgs]) if filtered_msgs else None
        if sel_msg_prop:
            idx_m = [f"{m['Предлагач']} ({m['Активни']} - {m['Брзина']})" for m in filtered_msgs].index(sel_msg_prop)
            chosen_m = filtered_msgs[idx_m]
            st.markdown(f"**Праќаме до Телеграм корисник:** `{chosen_m['Телеграм']}`")
            st.code(chosen_m['Порака'], language="text")

    with tg_tabs[1]:
        st.markdown("#### ✍️ Креирај и испрати твоја сопствена порака")
        st.info("Овде можеш да напишеш било каков текст (на пр. честитка, повик за состанок, итно известување) и да го пратиш до одредени тимови.")

        custom_msg = st.text_area(
            "Внеси го текстот на пораката:\n💡 Совет: Напиши **{име}** каде што сакаш да стои името на соговорникот (на пр. 'Здраво {име}, те молам...').",
            height=150
        )

        df_comp = tables_data['_comp']
        target_mode = st.radio(
            "Избери до кого да се прати:",
            ["Сите тимови", "Група IV (≥65%)", "Група III (50-65%)", "Група II (30-50%)", "Група I (≤30%)", "Рачен избор на тимови"],
            horizontal=True
        )

        selected_proposers = []
        if target_mode == "Сите тимови":
            selected_proposers = df_comp['Предлагач'].tolist()
        elif target_mode.startswith("Група"):
            gr_code = target_mode.split()[1] # IV, III, II, I
            selected_proposers = df_comp[df_comp['ГР.'] == gr_code]['Предлагач'].tolist()
        else:
            selected_proposers = st.multiselect("Избери предлагачи од листата:", df_comp['Предлагач'].tolist())

        st.write(f"📌 **Вкупно селектирани примачи:** {len(selected_proposers)}")

        if st.button("🔄 Генерирај листа за праќање (CSV)"):
            if not custom_msg.strip():
                st.error("Внеси текст на пораката!")
            elif not selected_proposers:
                st.error("Избери барем еден тим/предлагач!")
            else:
                custom_messages_data = []
                for prop in selected_proposers:
                    row = df_comp[df_comp['Предлагач'] == prop].iloc[0]
                    tg_n = row['Телеграм назив']
                    if pd.notna(tg_n) and str(tg_n).strip() != '':
                        first_name = str(prop).split()[0].split('(')[0].strip()
                        first_name = NAME_MAP_CYR.get(first_name, first_name)

                        final_msg = custom_msg.replace("{име}", first_name).replace("{ime}", first_name)

                        custom_messages_data.append({
                            'Предлагач': prop,
                            'Телеграм': str(tg_n).strip(),
                            'Порака': final_msg
                        })

                if custom_messages_data:
                    st.success(f"✅ Успешно генерирани пораки за {len(custom_messages_data)} корисници кои имаат Телеграм.")
                    df_custom = pd.DataFrame(custom_messages_data)
                    csv_data = df_custom.to_csv(index=False, encoding='utf-8-sig')
                    st.download_button(
                        label="📥 Преземи CSV за скриптата (prakac.py)",
                        data=csv_data,
                        file_name=f"telegram_custom_poraki.csv",
                        mime="text/csv",
                        type="primary",
                        key="btn_dl_custom_csv"
                    )
                    with st.expander("👁️ Преглед на генерираните пораки", expanded=False):
                        st.dataframe(df_custom, use_container_width=True)
                else:
                    st.warning("Избраните предлагачи немаат внесен Телеграм јузернејм во базата.")

elif menu.startswith("🔍 5."):
    st.markdown('<div class="sub-header">🔍 Брзо Пребарување низ Базата и Алијас Речник</div>', unsafe_allow_html=True)
    q = st.text_input("Внеси име, презиме, предлагач или алијас:")
    if q:
        res = df_main[
            df_main['Име и презиме / Профил'].astype(str).str.contains(q, case=False, na=False) |
            df_main['Предложувач'].astype(str).str.contains(q, case=False, na=False)
        ]
        st.info(f"Пронајдени се **{len(res)}** записи за '{q}':")
        st.dataframe(res[['Р.б.', 'Име и презиме / Профил', 'Предложувач', 'Име на Телеграм (За пребарување)'] + eval_cols[-3:]], use_container_width=True)

    with st.expander("📚 Преглед на вградениот Алијас Речник (Двојни профили, момински презимиња и страници)", expanded=True):
        alias_rows = [{"Име во База / Клуч": k, "Поврзани ФБ Профили / Алијаси": ", ".join(v)} for k, v in ALIAS_MAP.items()]
        st.dataframe(pd.DataFrame(alias_rows), use_container_width=True, height=400)

# ------------------------------------------------------------------------------
# ТАБ 6: МОБИЛЕН ПРИСТАП И ИКОНА НА ТЕЛЕФОН
# ------------------------------------------------------------------------------
elif menu.startswith("📱 6."):
    st.markdown('<div class="sub-header">📱 Поврзување од Мобилен Телефон и Додавање Икона</div>', unsafe_allow_html=True)
    try:
        hostname = socket.gethostname()
        local_ip = socket.gethostbyname(hostname)
    except Exception:
        local_ip = "192.168.1.X"

    st.markdown(f"""
    ### 1️⃣ Кога работиш преку Streamlit Cloud (Онлајн од секаде):
    * Откако ќе ги прикачиш **`app.py`**, **`{latest_file}`** и **`requirements.txt`** на GitHub и ќе ја пуштиш на [share.streamlit.io](https://share.streamlit.io), добиваш постојан линк.
    * Отвори го тој линк во Chrome (Android) или Safari (iPhone) и избери **„Add to Home Screen“ (Додади на почетен екран)** за да добиеш директна икона на телефонот.

    ### 2️⃣ Кога работиш локално преку ист Wi-Fi:
    * Стартувај во CMD со: `streamlit run app.py --server.address 0.0.0.0`
    * Отвори го овој линк на телефонот: **`http://{local_ip}:8501`**
    """)
